#!/usr/bin/env python3
"""Build the standalone Experiments, Results and Discussion manuscript section."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ANALYSIS = ROOT / "analysis"
TABLES = ANALYSIS / "derived" / "tables"
FIGURES = ANALYSIS / "figures"
VALIDATION = ANALYSIS / "validation"
DELIVERABLES = ROOT / "deliverables"
OUTPUT = DELIVERABLES / "Taxila_PreserveX_Experiments_Results_Discussion.docx"

# narrative_proposal preset, with the memo_masthead first-page pattern.
BLUE = "2E74B5"
DARK_BLUE = "16324F"
INK = "172B4D"
MUTED = "5B6573"
LIGHT = "F4F6F9"
PALE_BLUE = "EAF2F6"
AMBER = "FFF4D6"
AMBER_TEXT = "6F5100"
TEAL = "0F766E"
WHITE = "FFFFFF"
BORDER = "CDD5DF"


def set_run_font(run, name="Calibri", size=11, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("left", start), ("right", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in):
    widths_dxa = [round(width * 1440) for width in widths_in]
    if sum(widths_dxa) != 9360:
        widths_dxa[-1] += 9360 - sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        old_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_body(doc, text, keep=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.333
    paragraph.paragraph_format.keep_together = keep
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK)
    return paragraph


def add_equation(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, name="Cambria Math", size=10.5, color=DARK_BLUE, italic=True)
    return paragraph


def add_callout(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_together = True
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), AMBER)
    ppr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "5")
        border.set(qn("w:color"), "D9A441")
        p_bdr.append(border)
    ppr.append(p_bdr)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.2, bold=True, color=AMBER_TEXT)


def add_table_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=True, color=DARK_BLUE)


def add_source_note(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=8, italic=True, color=MUTED)


def add_table(doc, headers, rows, widths, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            if row_index == 0:
                shade(cell, LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        bold=(row_index == 0),
                        color=INK,
                    )
    return table


def add_figure(doc, number, filename, caption, alt_text, width=6.35):
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(6)
    image_paragraph.paragraph_format.space_after = Pt(2)
    image_paragraph.paragraph_format.keep_with_next = True
    run = image_paragraph.add_run()
    shape = run.add_picture(str(FIGURES / filename), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_paragraph.paragraph_format.keep_together = True
    caption_run = caption_paragraph.add_run(f"Figure {number}. {caption}")
    set_run_font(caption_run, size=8.5, italic=True, color=MUTED)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([start, instruction, end])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header.paragraphs[0]
    header.text = "Taxila PreserveX  |  Experiments, Results and Discussion"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def fmt(value, digits=3):
    return f"{float(value):.{digits}f}"


def pct(value, digits=2):
    return f"{100 * float(value):.{digits}f}%"


def build_document():
    DELIVERABLES.mkdir(parents=True, exist_ok=True)
    endpoint = json.loads((TABLES / "endpoint_screening_summary.json").read_text())
    climate = pd.read_csv(TABLES / "matched_epoch_climate_metrics.csv")
    scores = pd.read_csv(TABLES / "component_epoch_integrated_scores.csv")
    hotspots = pd.read_csv(TABLES / "component_endpoint_hotspots.csv")
    terrain = pd.read_csv(TABLES / "component_multiscale_terrain_hydrology.csv")
    models = pd.read_csv(TABLES / "spatially_blocked_model_comparison.csv")
    uncertainty = pd.read_csv(TABLES / "monte_carlo_domain_weight_rank_uncertainty.csv")
    ablation = pd.read_csv(TABLES / "leave_one_domain_factor_out_ablation.csv")
    scale = pd.read_csv(TABLES / "integrated_score_scale_rank_correlations.csv")
    associations = pd.read_csv(TABLES / "climate_spectral_exploratory_associations.csv")
    redundancy = pd.read_csv(TABLES / "factor_redundancy_screen.csv")
    validation = json.loads((VALIDATION / "integrated_experiment_validation.json").read_text())

    target = scores[(scores["epoch_id"] == "E2024") & (scores["radius_m"] == 500)].copy()
    target = target.sort_values("local_priority_score", ascending=False).reset_index(drop=True)
    target["integrated_rank"] = target.index + 1
    hot500 = hotspots[hotspots["radius_m"] == 500].copy()
    hot500 = hot500.sort_values("share_c_ge_2", ascending=False).reset_index(drop=True)
    hot500["spectral_rank"] = hot500.index + 1
    terrain500 = terrain[terrain["radius_m"] == 500].copy()
    terrain_scores = target[
        [
            "component_id",
            "terrain_susceptibility_score",
            "slope_pressure",
            "wetness_pressure",
            "drainage_proximity_pressure",
        ]
    ].merge(terrain500, on="component_id")
    terrain_scores = terrain_scores.sort_values("terrain_susceptibility_score", ascending=False)

    doc = Document()
    configure_styles(doc)
    configure_section(doc)

    # memo_masthead first-page pattern.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(6)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("STANDALONE MANUSCRIPT SECTION")
    set_run_font(run, size=9.5, bold=True, color=TEAL)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Experiments, Results and Discussion")
    set_run_font(run, size=23, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(
        "Climate-linked, multi-temporal assessment of landscape pressure around the Taxila World Heritage property, 2004–2024"
    )
    set_run_font(run, size=12, bold=True, color=MUTED)

    metadata = [
        ("Manuscript scope", "Experiments, Results and Discussion only"),
        ("Analytical units", "17 geolocated Taxila components; 250, 500 and 1,000 m supports"),
        ("Execution status", f"{validation['checks_passed']}/{validation['checks_total']} analytical validation gates passed"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=10, bold=True, color=INK)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

    add_callout(
        doc,
        "Executed headline findings: the endpoint screen retained 95.74% common support; 16.93% of supported cells met at least two adverse spectral criteria and 1.42% met all three. The 2024 climate window was dominated by unusually wet forcing. After hierarchical redundancy control, the Giri complex ranked first at 500 m, but its 95% weight-sensitive rank interval was 1–9. The MLP’s proxy macro-F1 was 0.887, yet tree-ensemble differences included zero under spatial-block bootstrap resampling.",
    )

    # Section 3: Experiments.
    doc.add_heading("3. Experiments", level=1)
    add_body(
        doc,
        "The experimental programme was designed to test one integrated scientific story without converting every available variable into a single opaque risk index. Six linked experiments were executed. The first reproduced the accepted Landsat endpoint screen on the fixed 30 m grid. The second matched daily weather observations to the five satellite epochs. The third quantified terrain and drainage-related susceptibility. The fourth compared transparent and nonlinear land-cover classifiers under spatial blocking. The fifth combined non-redundant pressure domains into relative component scores. The sixth tested whether the resulting priorities survived changes in spatial support, weights and factor composition. The experiments therefore move from observation to integration and finally to falsification of the ranking assumptions.",
    )
    add_table_caption(doc, "Table 1. Executed experiment matrix and permitted inference.")
    experiment_rows = [
        ("E1", "Endpoint reproduction", "Landsat 2004–2024", "Common support, adverse-tail convergence", "Spectral pressure"),
        ("E2", "Matched climate", "NASA POWER 1991–2025", "Wet, dry and heat extremes", "Site-wide temporal forcing"),
        ("E3", "Terrain/hydrology", "1 arc-second terrain tile", "Slope, TWI proxy, high-flow distance", "Static susceptibility"),
        ("E4", "Model comparison", "20,749 consensus samples", "Five blocked models + bootstrap", "Proxy agreement"),
        ("E5", "Integrated screen", "Landscape + terrain + climate", "Relative exposure and local priority", "Inspection priority"),
        ("E6", "Robustness", "Scales, ablations, 50,000 weights", "Rank stability and intervals", "Uncertainty bounds"),
    ]
    add_table(
        doc,
        ["ID", "Experiment", "Evidence", "Primary output", "Permitted inference"],
        experiment_rows,
        [0.45, 1.35, 1.45, 1.9, 1.35],
        font_size=7.6,
    )
    add_source_note(
        doc,
        "All experiments were executed from the frozen Landsat, component and land-cover bundles plus the cited NASA POWER and public terrain sources.",
    )

    doc.add_heading("3.1 Experiment 1: exact reproduction of the spectral endpoint screen", level=2)
    add_body(
        doc,
        "E1 reopened the five harmonised post-monsoon Landsat index composites for 2004, 2009, 2014, 2019 and 2024 and reconstructed the accepted 2004–2024 endpoint mask from the raster NoData structure. This was an integrity experiment before adding climate or terrain. The common grid comprised 403,480 cells in EPSG:32643. Only cells valid at both endpoints entered threshold derivation and convergence mapping. Adverse tails were frozen at the full common-support 20th percentiles of ΔNDVI and ΔMNDWI and the 80th percentile of ΔNDBI, preventing component-specific threshold tuning.",
    )
    add_body(
        doc,
        "The three binary criteria represented vegetation decline, wetness/surface-water decline and an increase in the built/bare spectral contrast. Their sum, C ∈ {0,1,2,3}, was used only as a spectral-convergence screen. A cell with C ≥ 2 was considered a multi-indicator hotspot, but not a damaged-monument observation. Component summaries were independently calculated at radii of 250, 500 and 1,000 m to expose scale dependence rather than treating 500 m as uniquely correct.",
    )

    doc.add_heading("3.2 Experiment 2: climate windows matched to satellite epochs", level=2)
    add_body(
        doc,
        "Daily NASA POWER records were obtained for the Taxila study location (33.746° N, 72.835° E) from 1991 through 2025. The parameters were precipitation, daily maximum and minimum temperature, mean temperature, relative humidity and 10 m wind speed. Each Landsat epoch was paired with a symmetric three-year weather window: 2003–2005, 2008–2010, 2013–2015, 2018–2020 and 2023–2025. This design reduces sensitivity to a single anomalous year while maintaining temporal correspondence with the satellite composites.",
    )
    add_body(
        doc,
        "A 1991–2020 reference climatology defined the 95th-percentile wet-day precipitation threshold (31.84 mm), daily maximum-temperature threshold (42.59 °C), diurnal-range threshold (18.17 °C) and strong-wind threshold (3.72 m s−1). Annual indicators included monsoon rainfall, maximum one- and five-day precipitation, very-wet-day count, consecutive wet and dry days, pre-composite 30- and 90-day rainfall, annual maximum temperature, hot-day count, heatwave duration, diurnal thermal range, relative humidity and strong-wind days. The composite climate-extreme score averaged baseline percentile ranks for RX5day, very-wet days, consecutive dry days, hot days, maximum heatwave duration and high-diurnal-range days.",
    )
    add_equation(
        doc,
        "Cₜ = (P_RX5 + P_R95 + P_CDD + P_TX95 + P_HW + P_DTR) / 6",
    )
    add_body(
        doc,
        "NASA POWER meteorology is much coarser than the 30 m Landsat grid. It was consequently treated as one site-wide forcing value for each epoch, not interpolated into monument-scale rainfall pixels. This distinction is central: climate can alter temporal exposure trajectories across the property, whereas within-epoch differences among components must be supplied by observed land-surface conditions and terrain susceptibility.",
    )

    doc.add_heading("3.3 Experiment 3: terrain and hydrological susceptibility", level=2)
    add_body(
        doc,
        "E3 used the public N33E072 one-arc-second bare-earth elevation tile and derived slope, terrain ruggedness, a 500 m topographic-position measure, D8 flow accumulation, a topographic-wetness proxy and distance to high-flow cells. D8 routing selected only a strictly lower neighbour; depressions and flats remained terminal rather than being artificially filled. The result is a reproducible flow-convergence susceptibility proxy, not a hydraulic inundation simulation. For each component and radius, the 90th-percentile slope and wetness proxy and median distance to high-flow cells were converted to pooled adverse-direction ranks and averaged into a terrain-susceptibility score.",
    )
    add_equation(
        doc,
        "Tᵢ,ᵣ = (R_slope + R_wetness + R_drainage-proximity) / 3",
    )

    doc.add_heading("3.4 Experiment 4: spatially blocked machine-learning comparison", level=2)
    add_body(
        doc,
        "Machine learning was included as a comparison experiment, not as decoration. The task was six-class discrimination against the existing WorldCover-consensus proxy label set. The feature table contained Landsat reflectance, spectral indices, support and spatial-context variables for 20,749 samples. Development used 16,554 samples in 79 spatial blocks, while the outer test retained 4,195 samples in 21 disjoint blocks; block overlap was exactly zero. Hyperparameters were selected inside the development data.",
    )
    add_body(
        doc,
        "Five models covered increasing nonlinear flexibility: multinomial logistic regression, random forest, extra trees, histogram gradient boosting and a multilayer perceptron with two hidden layers. Outer-test agreement, balanced accuracy, macro-F1 and log loss were calculated. To avoid treating neighbouring pixels as independent replicates, 2,000 bootstrap samples resampled entire held-out spatial blocks. The resulting macro-F1 intervals and pairwise differences were used to judge whether added model complexity delivered a material gain.",
    )
    add_body(
        doc,
        "The proxy labels are not substituted for the author-confirmed expert review of the land-cover reference sample. Because the locked per-unit expert-label file was not present in the executable bundle, the values reported in E4 are explicitly called proxy agreement, not independent thematic accuracy. Final expert-based user’s accuracy, producer’s accuracy and area-adjusted estimates should be inserted only from the locked expert file following the probability-sample protocol described by Olofsson et al. (2014).",
    )

    doc.add_heading("3.5 Experiment 5: non-redundant integrated component screen", level=2)
    add_body(
        doc,
        "All dynamic factors were transformed with pooled empirical cumulative distribution ranks so that scales remained comparable without implying a physical damage probability. A pre-integration Spearman screen detected near duplication between NDVI-pressure and NDBI-pressure (ρ = 0.972). Rather than allowing two correlated surface-cover variables to receive twice the weight of MNDWI, NDVI and NDBI were first averaged into one surface-cover subdomain. Mapped built-up and bare-class shares were retained for comparison but excluded from the primary composite because they are classifications derived from the same spectral evidence.",
    )
    add_equation(
        doc,
        "S_cover = (R_−NDVI + R_+NDBI) / 2;     L = (S_cover + R_−MNDWI) / 2",
    )
    add_body(
        doc,
        "The integrated exposure score E averaged landscape pressure L, terrain susceptibility T and climate forcing C. The local-priority score P averaged only L and T. Because C is constant across components within an epoch, E and P have the same within-epoch ordering; C instead changes the temporal trajectory between epochs. Reporting both quantities prevents coarse climate data from creating false monument-scale spatial precision.",
    )
    add_equation(
        doc,
        "Eᵢ,ₜ,ᵣ = (Lᵢ,ₜ,ᵣ + Tᵢ,ᵣ + Cₜ) / 3;     Pᵢ,ₜ,ᵣ = (Lᵢ,ₜ,ᵣ + Tᵢ,ᵣ) / 2",
    )

    doc.add_heading("3.6 Experiment 6: uncertainty, scale sensitivity and ablation", level=2)
    add_body(
        doc,
        "Three robustness tests challenged the primary 500 m ranking. First, 50,000 Dirichlet weight vectors were drawn for landscape, terrain and climate domains, and every component was re-ranked on each draw. The analysis reports median rank, 2.5th–97.5th percentile rank interval and probability of entering the top five or top three. Although climate weight changes absolute scores, climate is constant within 2024 and therefore cancels from the component ordering; rank uncertainty is driven by the balance between landscape and terrain evidence.",
    )
    add_body(
        doc,
        "Second, rank correlations were calculated among 250, 500 and 1,000 m supports. Third, the 500 m local-priority ranking was recomputed using landscape only, terrain only, and leave-one-factor-out variants. The hierarchical design was respected during ablation: removing NDVI or NDBI left the remaining surface-cover measure balanced against MNDWI, while removing MNDWI left the surface-cover subdomain balanced against terrain. These tests are not secondary decoration; they determine which component positions are stable enough to guide inspection and which must be presented as conditional.",
    )

    doc.add_heading("3.7 Reproducibility controls", level=2)
    add_body(
        doc,
        f"The final run wrote 22 machine-readable tables and nine publication figures and passed {validation['checks_passed']} of {validation['checks_total']} executable gates. The gates verified exact endpoint cell count, zero model-block overlap, finite bounded scores, expected tables and all figure files. SHA-256 hashes were recorded for NASA POWER data ({validation['source_hashes']['nasa_power_json'][:16]}…), the terrain tile ({validation['source_hashes']['terrain_hgt_gz'][:16]}…) and the raw raster manifest ({validation['source_hashes']['raw_manifest'][:16]}…). Random operations used a fixed seed. The complete code and validation JSON accompany this section.",
    )

    # Section 4: Results.
    doc.add_page_break()
    doc.add_heading("4. Results", level=1)
    doc.add_heading("4.1 Endpoint support and exact spectral reproduction", level=2)
    add_body(
        doc,
        f"The common endpoint mask contained {endpoint['endpoint_supported_cells']:,} cells, equivalent to {endpoint['endpoint_supported_percent']:.2f}% of the 403,480-cell analysis grid. The adverse thresholds were ΔNDVI ≤ {endpoint['thresholds']['delta_ndvi_q20']:.5f}, ΔMNDWI ≤ {endpoint['thresholds']['delta_mndwi_q20']:.5f} and ΔNDBI ≥ +{endpoint['thresholds']['delta_ndbi_q80']:.5f}. Of the supported cells, {endpoint['convergence_cell_counts']['0']:,} met none of the criteria, {endpoint['convergence_cell_counts']['1']:,} met one, {endpoint['convergence_cell_counts']['2']:,} met two and {endpoint['convergence_cell_counts']['3']:,} met all three. Thus, {endpoint['share_c_ge_2_percent']:.2f}% met at least two criteria and {endpoint['share_c_eq_3_percent']:.2f}% met all three.",
    )
    add_body(
        doc,
        f"Grid-wide median endpoint changes were +{endpoint['grid_median_delta']['NDVI']:.3f} for NDVI, {endpoint['grid_median_delta']['MNDWI']:.3f} for MNDWI and {endpoint['grid_median_delta']['NDBI']:.3f} for NDBI. The positive NDVI and negative NDBI medians show that the study grid did not undergo uniform adverse change. Hotspots are therefore tail events embedded within a landscape whose median surface-cover signal generally became greener and less built/bare-like. This is why a hotspot fraction cannot be interpreted as the proportion of the entire property that deteriorated.",
    )
    add_table_caption(doc, "Table 2. Reproduced endpoint-screen statistics.")
    endpoint_rows = [
        ("Common endpoint cells", f"{endpoint['endpoint_supported_cells']:,}", f"{endpoint['endpoint_supported_percent']:.2f}%"),
        ("ΔNDVI adverse threshold", f"{endpoint['thresholds']['delta_ndvi_q20']:.5f}", "20th percentile"),
        ("ΔMNDWI adverse threshold", f"{endpoint['thresholds']['delta_mndwi_q20']:.5f}", "20th percentile"),
        ("ΔNDBI adverse threshold", f"+{endpoint['thresholds']['delta_ndbi_q80']:.5f}", "80th percentile"),
        ("C ≥ 2", f"{endpoint['share_c_ge_2_percent']:.2f}%", "Multi-indicator spectral hotspot"),
        ("C = 3", f"{endpoint['share_c_eq_3_percent']:.2f}%", "All three spectral criteria"),
    ]
    add_table(doc, ["Metric", "Result", "Interpretation"], endpoint_rows, [2.05, 1.35, 3.1], font_size=8.5)

    doc.add_heading("4.2 Matched climate forcing", level=2)
    add_body(
        doc,
        "The five climate windows expressed different combinations of wet, dry and heat forcing rather than a monotonic trend. The 2024 window had the highest composite score (0.526), narrowly exceeding 2004 (0.517), followed by 2014 (0.484), 2009 (0.462) and 2019 (0.314). The low 2019 composite should not be read as a dry year: its mean annual rainfall was 1,918.5 mm and RX5day reached 237.0 mm, but consecutive dry days, hot days, heatwave duration and diurnal-range extremes were all comparatively low. The composite deliberately represents the simultaneous percentile position of several mechanisms.",
    )
    add_body(
        doc,
        "The 2024 window was distinctive for wet forcing. Mean monsoon rainfall reached 1,023.6 mm, 3.04 standard deviations above the 1991–2020 annual baseline distribution. RX5day was 231.6 mm (+1.27 SD) and the very-wet-day count averaged 10.0 (+1.17 SD). By contrast, annual maximum temperature, hot-day count, heatwave duration and mean diurnal range were below their baseline means. The correct interpretation is therefore a wet-extreme exposure window, not a compound wet-and-heat maximum.",
    )
    add_table_caption(doc, "Table 3. Matched climate-window results.")
    climate_rows = []
    for _, row in climate.iterrows():
        climate_rows.append(
            (
                int(row["label_year"]),
                str(row["years"]),
                f"{row['monsoon_rainfall_mm']:.1f}",
                f"{row['rx5day_mm']:.1f}",
                f"{row['r95p_days']:.1f}",
                f"{row['tx95_days']:.1f}",
                f"{row['climate_extreme_score']:.3f}",
            )
        )
    add_table(
        doc,
        ["Epoch", "Weather years", "Monsoon mm", "RX5day mm", "Very-wet days", "Hot days", "Climate score"],
        climate_rows,
        [0.55, 1.05, 1.0, 0.9, 1.0, 0.9, 1.1],
        font_size=7.6,
    )
    add_figure(
        doc,
        1,
        "figure_01_matched_climate_anomalies.png",
        "Standardised anomalies for eight climate indicators in the matched three-year windows. The 2024 window is dominated by wet anomalies, whereas thermal and diurnal-range anomalies are negative.",
        "Heatmap of climate anomalies for five Landsat epochs. The 2024 column shows strong positive monsoon rainfall, five-day rainfall and very-wet-day anomalies, with negative temperature and diurnal-range anomalies.",
        width=6.4,
    )
    add_figure(
        doc,
        2,
        "figure_02_climate_spectral_coevolution.png",
        "Site-wide climate-extreme scores and median Landsat indices. The figure supports temporal comparison but not causal attribution because only five unique epochs are available.",
        "Two-panel line chart with climate-extreme score above and median NDVI, MNDWI and NDBI below for 2004, 2009, 2014, 2019 and 2024.",
        width=6.2,
    )

    doc.add_heading("4.3 Terrain and drainage-related contrasts", level=2)
    add_body(
        doc,
        "Terrain separated steep, runoff-connected settings from low-slope wetness-convergent settings. At 500 m, the Giri Mosque and tombs had the largest median slope among the high-susceptibility group (15.38°; 90th percentile 29.54°), followed by Khader Mohra (14.55°) and Mohra Moradu (13.84°). The Giri complex combined a median slope of 12.89° with a short median distance to high-flow cells (138 m). Bhallar followed a different mechanism: its median slope was only 3.00°, but its wetness-pressure rank was 0.912 and median high-flow distance was 336 m. A single “high terrain score” therefore does not mean the same process at every component.",
    )
    add_table_caption(doc, "Table 4. Highest 500 m terrain-susceptibility scores and diagnostic metrics.")
    terrain_rows = []
    for _, row in terrain_scores.head(7).iterrows():
        terrain_rows.append(
            (
                row["component_name"],
                f"{row['terrain_susceptibility_score']:.3f}",
                f"{row['slope_median_deg']:.2f}",
                f"{row['slope_p90_deg']:.2f}",
                f"{row['twi_proxy_p90']:.2f}",
                f"{row['drainage_distance_median_m']:.0f}",
            )
        )
    add_table(
        doc,
        ["Component", "Terrain score", "Median slope", "P90 slope", "P90 TWI", "Flow distance m"],
        terrain_rows,
        [2.4, 0.8, 0.8, 0.75, 0.7, 1.05],
        font_size=7.6,
    )
    add_figure(
        doc,
        3,
        "figure_03_terrain_hydrology_atlas.png",
        "Terrain and hydrological-susceptibility atlas. White circles locate the georeferenced Taxila components; the four panels show elevation, slope, topographic wetness proxy and distance to high-flow cells.",
        "Four Taxila maps showing elevation, slope, topographic wetness proxy and distance to high-flow cells, with heritage component points overlaid.",
        width=6.35,
    )

    doc.add_heading("4.4 Factor redundancy and hierarchical weighting", level=2)
    add_body(
        doc,
        f"The redundancy screen identified one primitive-factor pair above |ρ| = 0.85: {redundancy.iloc[0]['factor_a']} versus {redundancy.iloc[0]['factor_b']} (ρ = {redundancy.iloc[0]['spearman_rho']:.3f}). The mapped built-up and bare shares also correlated with direct spectral factors, but they were not candidates for the primary composite. Hierarchical grouping changed the scientific interpretation of the 2024 ranking. It prevented a single surface-cover contrast from being counted twice and increased the relative influence of the independent wetness/surface-water subdomain.",
    )
    add_body(
        doc,
        "This correction materially changed which sites appeared most important. Bhallar remained the largest spectral-convergence hotspot but no longer ranked first in the integrated screen. That is a strength rather than a contradiction: the spectral map answers where adverse endpoint tails co-occur, while the integrated screen asks where non-redundant current pressure and terrain susceptibility converge. Publishing both makes the role of the index design visible.",
    )
    add_figure(
        doc,
        4,
        "figure_05_factor_redundancy_heatmap.png",
        "Spearman factor-correlation matrix. NDVI- and NDBI-derived pressures are grouped into one surface-cover subdomain before they are balanced against MNDWI.",
        "Correlation heatmap for NDVI pressure, MNDWI pressure, NDBI pressure, mapped built-up, mapped bare, the surface-cover subdomain and the final landscape score.",
        width=6.05,
    )

    doc.add_heading("4.5 Machine-learning comparison", level=2)
    add_body(
        doc,
        "The neural MLP returned the largest point macro-F1 (0.887), followed by extra trees (0.886), random forest (0.886) and histogram gradient boosting (0.885). Their overall proxy agreement ranged from 0.867 to 0.869. The point differences are operationally negligible. Whole-block resampling confirmed this: the extra-trees difference from the MLP had a 95% interval of −0.013 to +0.014; histogram gradient boosting, −0.012 to +0.014; and random forest, −0.014 to +0.012. The MLP was the bootstrap winner in only 41.9% of draws, compared with 28.1% for extra trees and 22.3% for histogram gradient boosting.",
    )
    add_body(
        doc,
        "Multinomial logistic regression was materially weaker, with macro-F1 = 0.863 and a difference from the MLP of −0.046 to −0.015. Thus, some nonlinearity was useful, but the experiment does not justify a claim that neural modelling outperformed tree ensembles. The broad MLP macro-F1 interval (0.666–0.906) also shows that performance varied substantially across the 21 outer spatial blocks. This spatial uncertainty should therefore accompany the point estimate rather than leaving the visually impressive 0.887 value to stand alone.",
    )
    add_table_caption(doc, "Table 5. Spatially blocked proxy-model comparison.")
    model_rows = []
    for _, row in models.iterrows():
        model_rows.append(
            (
                row["model"],
                f"{row['outer_proxy_overall_agreement']:.3f}",
                f"{row['outer_proxy_balanced_accuracy']:.3f}",
                f"{row['outer_proxy_macro_f1']:.3f}",
                f"[{row['outer_proxy_macro_f1_ci_low']:.3f}, {row['outer_proxy_macro_f1_ci_high']:.3f}]",
                f"{row['bootstrap_probability_best']:.3f}",
            )
        )
    add_table(
        doc,
        ["Model", "Agreement", "Balanced acc.", "Macro-F1", "95% block CI", "P(best)"],
        model_rows,
        [2.0, 0.8, 0.9, 0.75, 1.25, 0.8],
        font_size=7.7,
    )
    add_source_note(
        doc,
        "Metrics are held-out agreement with the WorldCover-consensus proxy under spatial blocking, not independent human-reference accuracy.",
    )
    add_figure(
        doc,
        5,
        "figure_06_spatially_blocked_model_comparison.png",
        "Five-model outer-test comparison and the MLP proxy confusion matrix. Macro-F1 whiskers are 95% whole-block bootstrap intervals; tree-ensemble differences from the MLP include zero.",
        "Horizontal bars compare macro-F1, balanced accuracy and overall agreement for five models with wide macro-F1 block-bootstrap error bars. A six-class MLP confusion matrix appears at right.",
        width=6.4,
    )

    doc.add_heading("4.6 Integrated exposure trajectories and 2024 component priorities", level=2)
    add_body(
        doc,
        "Integrated exposure was temporally non-monotonic. Scores generally declined in 2019 because the climate-extreme score was lowest and the pooled landscape-pressure ranks also shifted, then rebounded in 2024 with the wet-extreme forcing. This pattern argues against describing 2004–2024 as a single linear deterioration trajectory. The five epochs are more appropriately interpreted as repeated screening snapshots affected by changing climate forcing and land-surface context.",
    )
    add_figure(
        doc,
        6,
        "figure_04_component_exposure_trajectories.png",
        "Relative integrated exposure trajectories at the 500 m scale. Values are rank-based screening scores; the non-monotonic 2019 minimum and 2024 rebound are visible across components.",
        "Heatmap of 17 Taxila components by five epochs with integrated exposure scores. Most rows show lower values in 2019 and a partial rebound in 2024.",
        width=6.2,
    )
    add_body(
        doc,
        "At the primary 500 m scale, the highest 2024 local-priority score occurred at the Giri complex (0.591), followed by Giri Mosque and tombs (0.573), Jaulian (0.568), Mohra Moradu (0.539) and Dharmarajika (0.536). The Giri complex combined a landscape score of 0.506 with terrain susceptibility of 0.676. Giri Mosque and tombs had a similar combination (0.488 and 0.657), while Jaulian combined 0.518 and 0.618. These are balanced convergences rather than dominance by one domain.",
    )
    add_body(
        doc,
        "The spectral-only comparison produces a deliberately different list. Bhallar had the largest 500 m C ≥ 2 share (31.95%), followed by Jandial (24.88%), Khanpur Cave (23.31%), Pippala (23.17%) and Giri Mosque and tombs (23.11%). Under the integrated hierarchical score, Bhallar ranked eighth and Jandial sixteenth. The change means that their endpoint spectral tails did not coincide with equally high non-redundant 2024 landscape and terrain scores. This comparison prevents the manuscript from presenting a single index as an uncontested truth.",
    )
    union_names = list(dict.fromkeys(list(hot500.head(6)["component_name"]) + list(target.head(6)["component_name"])))
    compare = pd.DataFrame({"component_name": union_names})
    compare = compare.merge(
        hot500[["component_name", "spectral_rank", "share_c_ge_2"]],
        on="component_name",
        how="left",
    ).merge(
        target[["component_name", "integrated_rank", "local_priority_score"]],
        on="component_name",
        how="left",
    )
    compare_rows = []
    for _, row in compare.iterrows():
        compare_rows.append(
            (
                row["component_name"],
                int(row["spectral_rank"]) if pd.notna(row["spectral_rank"]) else "—",
                f"{100 * row['share_c_ge_2']:.2f}%" if pd.notna(row["share_c_ge_2"]) else "—",
                int(row["integrated_rank"]) if pd.notna(row["integrated_rank"]) else "—",
                f"{row['local_priority_score']:.3f}" if pd.notna(row["local_priority_score"]) else "—",
            )
        )
    add_table_caption(doc, "Table 6. Comparison of spectral-convergence and integrated 2024 priorities at 500 m.")
    add_table(
        doc,
        ["Component", "Spectral rank", "C ≥ 2 share", "Integrated rank", "Local score"],
        compare_rows,
        [2.65, 0.85, 0.85, 0.95, 1.2],
        font_size=7.5,
    )
    add_figure(
        doc,
        7,
        "figure_07_integrated_hotspot_map.png",
        "Endpoint spectral convergence and 2024 hierarchical local priority. The left map retains the observed C = 0–3 raster; point colour and the right panel show the non-redundant component score.",
        "Map of Taxila spectral convergence with coloured component points, paired with a horizontal bar chart led by Giri complex, Giri Mosque and tombs, and Jaulian.",
        width=6.4,
    )

    doc.add_heading("4.7 Weight uncertainty, scale sensitivity and ablation", level=2)
    add_body(
        doc,
        "The weight experiment rejected the idea of a fully stable rank order. Giri had the highest equal-domain integrated score and a median Monte Carlo rank of 1, but its 95% interval extended from 1 to 9 and its top-five probability was 83.2%. Giri Mosque and tombs had a 2–11 interval and 73.0% top-five probability; Jaulian had a narrower 2–7 interval and the largest top-five probability among the leaders (83.4%). Bhallar’s equal-domain rank was eighth, yet its interval spanned 1–15 because it moves sharply upward when terrain receives more weight. Sirkap showed a similarly broad 1–15 interval because its high landscape score is offset by lower terrain susceptibility.",
    )
    add_table_caption(doc, "Table 7. Highest equal-domain integrated scores with weight-sensitive rank uncertainty.")
    uncertainty_rows = []
    for _, row in uncertainty.head(10).iterrows():
        uncertainty_rows.append(
            (
                row["component_name"],
                int(row["equal_domain_rank"]),
                int(row["monte_carlo_median_rank"]),
                f"{int(row['rank_p2_5'])}–{int(row['rank_p97_5'])}",
                f"{100 * row['probability_top_5']:.1f}%",
                f"{100 * row['probability_top_3']:.1f}%",
            )
        )
    add_table(
        doc,
        ["Component", "Equal rank", "Median rank", "95% rank interval", "P(top 5)", "P(top 3)"],
        uncertainty_rows,
        [2.65, 0.65, 0.75, 1.0, 0.7, 0.75],
        font_size=7.5,
    )
    add_figure(
        doc,
        8,
        "figure_08_monte_carlo_rank_uncertainty.png",
        "Component rank intervals from 50,000 domain-weight draws. Wide intervals show why equal-weight ranks should be used as screening priorities rather than definitive risk positions.",
        "Dot-and-whisker chart of 17 component ranks. Giri is centered near rank 1 but extends to 9; Bhallar and Sirkap have especially wide intervals.",
        width=6.25,
    )
    rho_250_500 = scale[(scale["radius_a_m"] == 250) & (scale["radius_b_m"] == 500)]["spearman_rank_correlation"].iloc[0]
    rho_500_1000 = scale[(scale["radius_a_m"] == 500) & (scale["radius_b_m"] == 1000)]["spearman_rank_correlation"].iloc[0]
    rho_250_1000 = scale[(scale["radius_a_m"] == 250) & (scale["radius_b_m"] == 1000)]["spearman_rank_correlation"].iloc[0]
    add_body(
        doc,
        f"Scale agreement was strongest between 500 and 1,000 m (ρ = {rho_500_1000:.3f}), followed by 250 and 500 m (ρ = {rho_250_500:.3f}) and 250 and 1,000 m (ρ = {rho_250_1000:.3f}). Giri ranked fourth, first and second across the three supports; Jaulian ranked second, third and first; and Giri Mosque and tombs ranked first, second and fourth. These three components are therefore persistently high even though their exact order is scale-dependent.",
    )
    add_body(
        doc,
        "Ablation exposed the ranking’s mechanism. Terrain-only agreement with the full local ranking was ρ = 0.819 with four of five leaders retained, whereas landscape-only agreement was ρ = 0.316 with only one leader retained. Removing NDVI or NDBI individually barely changed the hierarchical ranking (ρ = 0.988 and 0.990), confirming that grouping successfully prevented either correlated surface-cover factor from controlling the result. Removing MNDWI reversed much of the order (ρ = −0.275; zero top-five overlap). The main conclusion is not that MNDWI is unquestionably correct, but that the final priority list is conditional on the wetness/surface-water pressure mechanism and must be checked with field moisture, drainage and seasonal-water evidence.",
    )
    stability_rows = [
        ("250 vs 500 m", f"{rho_250_500:.3f}", "Moderate–high"),
        ("500 vs 1,000 m", f"{rho_500_1000:.3f}", "High"),
        ("250 vs 1,000 m", f"{rho_250_1000:.3f}", "Moderate"),
    ]
    for _, row in ablation.iterrows():
        stability_rows.append(
            (
                row["ablation"],
                f"{row['spearman_rho_with_full']:.3f}",
                f"Top-five overlap {int(row['top5_overlap_with_full'])}/5",
            )
        )
    add_table_caption(doc, "Table 8. Scale and ablation stability.")
    add_table(doc, ["Comparison", "Spearman ρ", "Interpretation"], stability_rows, [2.55, 1.0, 2.95], font_size=8.0)
    add_figure(
        doc,
        9,
        "figure_09_scale_and_ablation_sensitivity.png",
        "Integrated ranks by spatial support and leave-one-domain/factor-out stability. MNDWI removal produces the strongest departure from the full ranking.",
        "Two-panel figure with a heatmap of ranks at 250, 500 and 1,000 m and bars of Spearman stability for full, terrain-only, landscape-only and factor-removal variants.",
        width=6.35,
    )

    doc.add_heading("4.8 Exploratory climate–spectral co-variation", level=2)
    add_body(
        doc,
        "With only five unique epochs, climate–spectral associations were treated as effect-direction screens without confirmatory p-values. Monsoon rainfall and median NDVI had ρ = 0.900, with the leave-one-epoch-out correlation remaining between 0.800 and 1.000. Hot-day count and NDVI had ρ = −0.600, stable between −0.800 and −0.400. The composite climate score and MNDWI had ρ = −0.600 with the same sign in every leave-one-out fit. These patterns are compatible with climate-linked surface response, but they cannot distinguish direct climatic effects from land management, sensor residuals or coincident land-cover change.",
    )
    add_table_caption(doc, "Table 9. Exploratory climate–spectral associations across five epochs.")
    association_rows = []
    for _, row in associations.iterrows():
        association_rows.append(
            (
                row["climate_metric"].replace("_", " "),
                row["spectral_metric"].replace("grid_median_", "").upper(),
                f"{row['spearman_rho_n5']:.2f}",
                f"{row['loo_rho_min']:.2f} to {row['loo_rho_max']:.2f}",
                f"{row['loo_sign_stability']:.2f}",
            )
        )
    add_table(
        doc,
        ["Climate factor", "Spectral metric", "ρ (n=5)", "Leave-one-out range", "Sign stability"],
        association_rows,
        [2.25, 1.05, 0.75, 1.45, 1.0],
        font_size=7.7,
    )

    # Section 5: Discussion.
    doc.add_page_break()
    doc.add_heading("5. Discussion", level=1)
    doc.add_heading("5.1 Main contribution of the integrated experiments", level=2)
    add_body(
        doc,
        "The experiments transform the paper from a three-index hotspot exercise into a climate-linked, terrain-aware and uncertainty-explicit screening study. The strongest contribution is not the production of a larger composite score. It is the separation of four distinct evidentiary questions: where adverse spectral endpoint tails co-occur; when the wider property experienced unusual wet, dry or heat forcing; where terrain can amplify runoff, wetness or erosion susceptibility; and how sensitive component priorities are to modelling choices. Each question has a corresponding map or statistic, and no layer is allowed to claim monument damage by itself.",
    )
    add_body(
        doc,
        "This framing complements earlier Taxila work that combined ground-based and satellite techniques for risk evaluation (Khan et al., 2022). The present study adds a fixed-grid multi-epoch design, climate windows matched to satellite observations, hierarchical redundancy control, spatially blocked algorithm comparison and explicit rank uncertainty. Those additions are particularly relevant for a serial World Heritage property, where visually compelling maps can otherwise conceal mismatched spatial resolution, duplicated factors and unstable rankings.",
    )

    doc.add_heading("5.2 What the climate experiment adds—and what it cannot add", level=2)
    add_body(
        doc,
        "The climate results supply plausible temporal forcing for the observed landscape signals. The extreme wetness of the 2024 window is relevant to surface erosion, runoff concentration, biological activity and moisture cycling, and it provides a defensible reason to prioritise post-monsoon inspection. However, NASA POWER’s meteorological grid is approximately 0.5° × 0.625°. Its strength here is daily temporal consistency over 1991–2025, not spatial discrimination among monuments separated by only a few kilometres. Treating the climate score as constant within each epoch avoids false precision and follows the provider’s guidance not to request or represent the data at finer resolution than its source grid.",
    )
    add_body(
        doc,
        "The exploratory association between monsoon rainfall and NDVI is strong in rank terms, but five observations cannot establish a process model. A wet window can increase vegetation greenness while also intensifying erosion or moisture stress at exposed masonry; these responses are not contradictory. Future field linkage should therefore test mechanism-specific outcomes such as dampness, salt activity, rilling, surface loss or drainage failure rather than correlating the climate score with one generic condition label.",
    )

    doc.add_heading("5.3 Why the spectral and integrated rankings differ", level=2)
    add_body(
        doc,
        "The spectral-convergence map and integrated priority table answer different questions. Bhallar’s 31.95% C ≥ 2 share identifies extensive endpoint-tail overlap in its 500 m neighbourhood. Its eighth-place integrated position indicates that, after correlated surface-cover signals are grouped and terrain is balanced with landscape pressure, the complete local evidence is less dominant than at Giri, Giri Mosque and tombs or Jaulian. Jandial’s decline from second in spectral convergence to sixteenth in the integrated screen is even more pronounced. These differences should be presented side by side, because they help conservation teams choose the reason for inspection: rapid land-surface change, terrain-amplified susceptibility, or both.",
    )
    add_body(
        doc,
        "The hierarchical correction is scientifically consequential. Before correction, correlated NDVI and NDBI pressures would have received two-thirds of landscape weight and Bhallar would have ranked first. After they were treated as one surface-cover subdomain, the Giri settings rose because their MNDWI-derived pressure and terrain susceptibility converged. This result demonstrates why factor screening must occur before rather than after a hotspot map is published. It also explains the strong MNDWI ablation effect: the independent wetness/surface-water subdomain carries information not reproduced by the highly correlated surface-cover pair.",
    )

    doc.add_heading("5.4 Interpretation of the machine-learning comparison", level=2)
    add_body(
        doc,
        "The model experiment supports a restrained conclusion. Nonlinear models outperformed the multinomial baseline, but the MLP did not materially outperform extra trees, random forest or histogram gradient boosting. A deeper neural architecture would add parameters and tuning burden without evidence of a generalisation benefit, particularly because the outer evaluation contains only 21 spatial blocks. For the final paper, extra trees or the MLP may be retained as the best proxy models, but the scientific emphasis should remain on blocked validation and uncertainty rather than on the prestige of the algorithm name.",
    )
    add_body(
        doc,
        "The confusion matrix shows that water was perfectly separated in the proxy test, whereas confusion remained among built, bare, crop, woody and shrub/grass classes. These are precisely the boundaries that can affect anthropogenic-pressure estimates. Once the locked expert labels are available, model selection should be repeated or at least recalibrated against that independent reference. Until then, the current scores demonstrate discriminative plausibility and comparative model behaviour, not final map accuracy.",
    )

    doc.add_heading("5.5 Robust priorities and conditional priorities", level=2)
    add_body(
        doc,
        "Giri complex, Giri Mosque and tombs, and Jaulian remain near the top across spatial supports and retain high top-five probabilities under weight variation. They form the most defensible first tier for integrated field inspection. Dharmarajika and Mohra Moradu form a second tier with high equal-weight positions but broader uncertainty. Bhallar and Sirkap are conditional priorities: each can become first under plausible domain weights, but each also falls well outside the top five under other weights. Conservation reporting should therefore use priority tiers and rank intervals, not a deterministic numbered league table.",
    )
    add_body(
        doc,
        "The scale experiment also clarifies operational use. The high 500–1,000 m agreement suggests that broader landscape context is reasonably stable, whereas the weaker 250–1,000 m relationship indicates that immediate surroundings can tell a different story. A practical inspection protocol could use the 250 m map to identify direct encroachment or drainage concerns and the 1,000 m map to diagnose catchment-scale or access-related pressure, with 500 m serving as the comparative primary scale rather than the only valid scale.",
    )

    doc.add_heading("5.6 Conservation implications for Taxila", level=2)
    add_body(
        doc,
        "The integrated results support differentiated inspection rather than one uniform response. At the Giri complex and Giri Mosque and tombs, steep slopes and short distance to high-flow cells justify checking runoff paths, slope instability, drainage maintenance and erosion after intense monsoon periods. At Jaulian and Mohra Moradu, the combination of slope and drainage proximity warrants similar attention, while the spectral subdomains can guide inspection of vegetation, exposed soil and new built/bare surfaces. At Bhallar, the combination of extensive spectral convergence and a wetness-dominated terrain score suggests that hydrological context should be investigated even though the balanced integrated rank is lower.",
    )
    add_body(
        doc,
        "These recommendations remain triage actions. The satellite and terrain evidence do not reveal cracking, stone detachment, salt crystallisation, mortar loss or foundation movement. Field teams should therefore record those condition variables separately and link them to the pressure domains. A future management-priority model can then test whether the present exposure scores explain independently observed condition severity, rather than assuming that an exposure map is already a validated risk map.",
    )

    doc.add_heading("5.7 Limitations", level=2)
    add_body(
        doc,
        "Four limitations bound the conclusions. First, the climate record is temporally rich but spatially coarse; it cannot rank monuments within one epoch. Second, the terrain flow analysis is a topographic proxy with terminal sinks and no rainfall–runoff calibration, channel survey or engineered-drainage representation. Third, five Landsat epochs are adequate for matched maps but insufficient for confirmatory time-series causality. Annual component series would be needed for robust lagged climate analysis. Fourth, the integrated scores are rank-based and relative to the present components, epochs and supports. Adding a new component or factor can change the empirical ranks.",
    )
    add_body(
        doc,
        "The validation boundary is equally important. The author has confirmed independent domain-expert review of the land-cover reference interpretation, but the locked per-unit expert file was not available in the executable bundle used here. The model comparison is therefore reported against a consensus proxy. No pressure score in this section has been validated against a monument-condition outcome. These constraints do not invalidate the screening analysis; they define which claims belong in the paper now and which require the final expert and condition datasets.",
    )

    doc.add_heading("5.8 Implications for the paper’s central claim", level=2)
    add_body(
        doc,
        "The evidence supports a clear, defensible claim: PreserveX provides a reproducible way to integrate multi-temporal satellite change, matched weather extremes and terrain susceptibility while exposing redundancy, spatial-scale effects, model uncertainty and weight sensitivity. It identifies places and mechanisms that merit field inspection. The evidence does not support claiming automated monument-damage detection, a calibrated probability of heritage loss or a universally correct site ranking.",
    )
    add_callout(
        doc,
        "Synthesis: the Giri settings and Jaulian are robust first-tier integrated inspection priorities; Bhallar remains the strongest endpoint spectral hotspot; and all component positions should be communicated with scale and weight uncertainty. The study’s novelty lies in auditable evidence integration and uncertainty governance, not in presenting one final black-box risk score.",
    )

    doc.add_heading("References cited in Sections 3–5", level=1)
    references = [
        "Khan, M. Y., et al. (2022). Evaluation of risks to UNESCO World Heritage sites: Taxila, Pakistan using ground-based and satellite remote sensing techniques. Journal of Cultural Heritage, 55, 195–209. https://doi.org/10.1016/j.culher.2022.03.011",
        "NASA POWER. Daily API documentation and parameter specifications. https://power.larc.nasa.gov/docs/services/api/temporal/daily/",
        "NASA POWER. API data-request and meteorological-resolution guidance. https://power.larc.nasa.gov/docs/tutorials/service-data-request/api/",
        "Olofsson, P., Foody, G. M., Herold, M., Stehman, S. V., Woodcock, C. E., & Wulder, M. A. (2014). Good practices for estimating area and assessing accuracy of land change. Remote Sensing of Environment, 148, 42–57. https://doi.org/10.1016/j.rse.2014.02.015",
        "Registry of Open Data on AWS. Terrain Tiles: global bare-earth terrain heights. https://registry.opendata.aws/terrain-tiles/",
    ]
    for reference in references:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(reference)
        set_run_font(run, size=9, color=INK)

    doc.core_properties.title = "Taxila PreserveX Experiments, Results and Discussion"
    doc.core_properties.subject = "Climate-linked, terrain-aware and uncertainty-explicit experiments for the Taxila World Heritage property"
    doc.core_properties.author = "PreserveX research workflow"
    doc.core_properties.keywords = "Taxila; Landsat; NASA POWER; climate extremes; terrain; machine learning; uncertainty; cultural heritage"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
